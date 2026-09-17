"""Word findings report (python-docx) to go with the Excel workbook.

The workbook is for the person checking the evidence; this is for the board
or a funder: an executive summary in sentences, the profile by function and
by category, and the actions to start with.
"""

from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .engine import Engine, Result
from .heatmap import NOT_APPLICABLE, attainment_colour
from .rollup import by_category, by_function

TOP_ACTIONS = 10
PROFILE_HEADERS = ("Function", "Subcategories assessed", "At target", "Current (0-3)", "Target (0-3)", "Target reached")
CATEGORY_HEADERS = ("Category", "Current", "Target", "Reached", "Worst gap rating")
ACTION_HEADERS = ("Rank", "CSF ID", "What to do", "Cost", "Risk now", "Reduced by")
RATING_COLOURS = {"Low": "C6EFCE", "Medium": "FFEB9C", "High": "F8CBAD", "Critical": "FF7C80"}


def _shade(cell, hex_colour: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour.lstrip("#").upper())
    cell._tc.get_or_add_tcPr().append(shd)


def _table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, "1F3864")
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = "" if value is None else str(value)
    if widths:
        table.autofit = False
        for col, w in zip(table.columns, widths):
            col.width = Cm(w)
        for row in table.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Cm(w)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return table


def _pct(x: float | None) -> str:
    return "n/a" if x is None else f"{round(100 * x)}%"


def summary_points(engine: Engine, result: Result, organisation: str) -> list[str]:
    """The executive summary as plain sentences (also used by the web page)."""
    t = result.totals
    functions = [f for f in by_function(engine.catalog, result) if f.attainment is not None]
    overall_reached = sum(min(r.current, r.target) for r in result.subcategories if r.status == "assessed")
    overall_wanted = sum(r.target for r in result.subcategories if r.status == "assessed")
    weakest = min(functions, key=lambda f: f.attainment)
    strongest = max(functions, key=lambda f: f.attainment)
    rated = [r for r in result.subcategories if r.status == "assessed" and r.gap > 0]
    severe = sum(r.rating in ("High", "Critical") for r in rated)
    free_now = sum(a.cost_band == "free" for a in result.actions[:TOP_ACTIONS])
    points = [
        f"{organisation} answered {t['questions_asked']} of the {t['questions_in_bank']} questions in the bank; "
        f"the other {t['questions_in_bank'] - t['questions_asked']} did not apply to its profile or were "
        f"follow-ups to practices it does not have yet.",
        f"Of the {t['subcategories_assessed']} CSF 2.0 subcategories that apply, {t['subcategories_at_target']} "
        f"meet the target taken from NIST's Small Business Quick-Start Guide. Overall the organisation reaches "
        f"{_pct(overall_reached / overall_wanted)} of its target profile.",
        f"The strongest function is {strongest.title.capitalize()} ({_pct(strongest.attainment)} of target) "
        f"and the weakest is {weakest.title.capitalize()} ({_pct(weakest.attainment)}).",
        f"{len(rated)} gaps remain, {severe} of them rated High or Critical for likelihood and impact.",
    ]
    if result.actions:
        points.append(
            f"The action plan has {len(result.actions)} actions ranked by risk reduced per unit of cost; "
            f"{free_now} of the top {min(TOP_ACTIONS, len(result.actions))} need staff time rather than money."
        )
    if not result.complete:
        points.append("The assessment is not finished: unanswered questions were scored as 0.")
    return points


def build_report(
    engine: Engine,
    result: Result,
    answers: Mapping[str, str],
    evidence: Mapping[str, str] | None = None,
    organisation: str = "Organisation",
    top: int = TOP_ACTIONS,
):
    evidence = evidence or {}
    cat = engine.catalog
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)

    doc.add_heading("Cybersecurity self-assessment findings", level=0)
    p = doc.add_paragraph()
    p.add_run(organisation).bold = True
    p.add_run(f"\nAssessed against {cat.document} {cat.version}. Targets from {engine.targets.source}.")

    doc.add_heading("Executive summary", level=1)
    for point in summary_points(engine, result, organisation):
        doc.add_paragraph(point, style="List Bullet")
    if result.actions:
        doc.add_paragraph("Start with these:")
        for a in result.actions[:5]:
            doc.add_paragraph(f"{a.what_to_do} ({a.subcategory}; {a.cost_label})", style="List Number")

    doc.add_heading("Profile by function", level=1)
    rows = []
    for f in by_function(cat, result):
        if f.current is None:
            rows.append((f.title.capitalize(), 0, 0, None, None, "n/a"))
        else:
            rows.append((f.title.capitalize(), f.assessed, f.at_target, f"{f.current:.2f}", f"{f.target:.2f}",
                         _pct(f.attainment)))
    _table(doc, PROFILE_HEADERS, rows, widths=(3.5, 3, 2, 2.5, 2.5, 2.5))

    doc.add_heading("Profile by category", level=1)
    doc.add_paragraph(
        "Current and target are averages on the 0-3 scale over the subcategories that apply. "
        "Shading follows the heat map: red where little of the target is reached, green where all of it is."
    )
    grid = by_category(cat, result)
    flat = [c for cells in grid.values() for c in cells]
    rows = [
        (f"{c.id} {c.title}",
         None if c.current is None else f"{c.current:.2f}",
         None if c.target is None else f"{c.target:.2f}",
         _pct(c.attainment),
         c.worst_rating or ("not applicable" if c.current is None else "none"))
        for c in flat
    ]
    table = _table(doc, CATEGORY_HEADERS, rows, widths=(8, 2, 2, 2, 3))
    for row, c in zip(table.rows[1:], flat):
        _shade(row.cells[3], attainment_colour(c.attainment) if c.attainment is not None else NOT_APPLICABLE)
        if c.worst_rating:
            _shade(row.cells[4], RATING_COLOURS[c.worst_rating])

    doc.add_heading(f"Top {min(top, len(result.actions))} actions", level=1)
    doc.add_paragraph(
        "Each action closes one gap. The wording is NIST's implementation example for the subcategory; "
        "the full plan with owners and status columns is in the Excel workbook."
    )
    _table(
        doc, ACTION_HEADERS,
        [(a.rank, a.subcategory, a.what_to_do, a.cost_label, a.risk, a.risk_reduction) for a in result.actions[:top]],
        widths=(1.3, 2.3, 7.4, 2.8, 1.5, 1.7),
    )

    excluded = [r for r in result.subcategories if r.status == "excluded"]
    if excluded:
        doc.add_heading("Not applicable", level=1)
        reasons: dict[str, list[str]] = {}
        for r in excluded:
            reasons.setdefault(r.reason, []).append(r.id)
        for reason, ids in reasons.items():
            doc.add_paragraph(f"{', '.join(ids)}: {reason}", style="List Bullet")

    doc.add_heading("About this assessment", level=1)
    with_evidence = sum(1 for q in answers if evidence.get(q))
    doc.add_paragraph(
        f"This is a self-assessment: it records what the organisation says it does. No systems were scanned "
        f"and it is not an ISO 27001 or SOC 2 certification. Evidence was offered for {with_evidence} of "
        f"{len(answers)} answers; the Evidence Log sheet of the workbook is where a reviewer records what "
        f"they actually saw. Likelihood and impact values are coarse judgements for a small non-profit."
    )
    return doc


def export_report(path: Path, *args, **kwargs) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_report(*args, **kwargs).save(path)
    return path
