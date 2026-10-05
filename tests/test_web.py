"""The server-rendered questionnaire, walked through like a person would."""

import io
import json
import re

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from csfcheck.api import create_app
from csfcheck.engine import ASKED

from .conftest import SAMPLES

RADIO = re.compile(r'name="q_([a-z0-9-]+)" value="([^"]+)"')


@pytest.fixture(scope="module")
def app(database_url):
    app = create_app(database_url)
    yield app
    app.state.db.dispose()


@pytest.fixture
def client(app):
    with TestClient(app) as c:
        yield c


def start(client, name, profile):
    form = {"name": name, **{flag: "on" for flag, on in profile.items() if on}}
    r = client.post("/start", data=form)
    assert r.status_code == 200, r.text
    aid = int(re.search(r"/a/(\d+)", str(r.url)).group(1))
    return aid, r


def shown_questions(html):
    return list(dict.fromkeys(q for q, _ in RADIO.findall(html)))


def test_start_page_lists_profile_flags(client, engine):
    r = client.get("/")
    assert r.status_code == 200
    for flag in engine.bank.flags:
        assert f'name="{flag}"' in r.text


def test_start_needs_a_name(client):
    r = client.post("/start", data={"name": "  "})
    assert r.status_code == 422 and "give the organisation a name" in r.text


@pytest.mark.parametrize("path", SAMPLES, ids=[p.stem for p in SAMPLES])
def test_sample_walkthrough(client, engine, path):
    sample = json.loads(path.read_text())
    org, answers, evidence = sample["organisation"], sample["answers"], sample["evidence"]
    aid, r = start(client, org["name"], org["profile"])
    assert "resume code" in r.text

    final = engine.question_status(org["profile"], answers)
    seen: list[str] = []
    r = client.get(f"/a/{aid}/questions")
    for _ in range(40):
        if r.url.path.endswith("/heatmap"):
            break
        assert re.search(r"/a/\d+/s/[a-z-]+$", r.url.path), r.url
        qs = shown_questions(r.text)
        assert qs, f"empty page at {r.url}"
        # Only questions that apply to this profile are put in front of the user.
        assert all(final[q] == ASKED for q in qs)
        assert "questions answered" in r.text
        seen += [q for q in qs if q not in seen]
        form = {"go": "next"}
        for q in qs:
            form[f"q_{q}"] = answers[q]
            form[f"e_{q}"] = evidence.get(q, "")
        r = client.post(str(r.url.path), data=form)
        assert r.status_code == 200, r.text
    else:
        pytest.fail("questionnaire did not finish")

    assert set(seen) == set(answers)
    result = engine.assess(org["profile"], answers)
    t = result.totals
    page = r.text
    assert f"{len(answers)} of {len(answers)} questions answered" in page
    assert page.count('class="cell"') == 22
    assert f"{t['subcategories_at_target']} meet the target" in page

    plan = client.get(f"/a/{aid}/plan")
    assert plan.status_code == 200
    ranks = [int(x) for x in re.findall(r"<tr>\s*<td>(\d+)</td>", plan.text)]
    assert ranks == list(range(1, len(result.actions) + 1))
    assert result.actions[0].what_to_do.replace("'", "&#39;") in plan.text

    svg = client.get(f"/a/{aid}/heatmap.svg")
    assert svg.headers["content-type"].startswith("image/svg+xml")

    doc = Document(io.BytesIO(client.get(f"/a/{aid}/report.docx").content))
    assert org["name"] in doc.paragraphs[1].text
    wb = load_workbook(io.BytesIO(client.get(f"/a/{aid}/workbook.xlsx").content))
    log = {row[0].value: row[6].value for row in wb["Evidence Log"].iter_rows(min_row=2)}
    assert set(log) == set(answers)
    for q, note in evidence.items():
        assert log[q] == note


def test_follow_up_opens_on_the_same_page(client):
    aid, _ = start(client, "Backups only", {})
    r = client.get(f"/a/{aid}/s/protect")
    assert "pr-backup-test" not in shown_questions(r.text)
    assert "A follow-up question appears here" in r.text
    r = client.post(f"/a/{aid}/s/protect", data={"q_pr-backup": "yes"}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].endswith("/s/protect#q-pr-backup-test")
    assert "pr-backup-test" in shown_questions(client.get(r.headers["location"]).text)


def test_lowering_a_parent_closes_its_follow_up(client):
    aid, _ = start(client, "Backups lapsed", {})
    client.post(f"/a/{aid}/s/protect", data={"q_pr-backup": "yes"})
    client.post(f"/a/{aid}/s/protect", data={"q_pr-backup": "yes", "q_pr-backup-test": "yes"})
    # The follow-up's old answer is still in the form, as a browser would send it.
    r = client.post(f"/a/{aid}/s/protect", data={"q_pr-backup": "no", "q_pr-backup-test": "yes"},
                    follow_redirects=False)
    assert r.status_code == 303
    page = client.get(f"/a/{aid}/s/protect").text
    assert "pr-backup-test" not in shown_questions(page)
    assert "Skipped because you answered" in page


def test_profile_skipped_questions_are_explained(client):
    aid, _ = start(client, "No premises", {})
    page = client.get(f"/a/{aid}/s/protect").text
    assert "pr-physical" not in shown_questions(page)
    assert "Only asked if: has office or has servers" in page


def test_save_and_go_back(client):
    aid, _ = start(client, "Back button", {})
    r = client.post(f"/a/{aid}/s/identify", data={"go": "back"}, follow_redirects=False)
    assert r.headers["location"].endswith(f"/a/{aid}/s/govern")


def test_invalid_answer_is_rejected(client):
    aid, _ = start(client, "Bad value", {})
    r = client.post(f"/a/{aid}/s/govern", data={"q_gv-mission": "perhaps"})
    assert r.status_code == 422


def test_pages_need_the_assessment_cookie(app, client):
    aid, r = start(client, "Cookie check", {"has_office": True})
    token = re.search(r'<code class="token">([^<]+)</code>', r.text).group(1)
    with TestClient(app) as other:
        assert other.get(f"/a/{aid}/heatmap").status_code == 404
        assert other.get(f"/a/{aid}/report.docx").status_code == 404
        bad = other.post("/resume", data={"assessment_id": aid, "token": "wrong"})
        assert bad.status_code == 404 and "No assessment matches" in bad.text
        ok = other.post("/resume", data={"assessment_id": aid, "token": token})
        assert ok.status_code == 200 and "Cookie check" in ok.text
        assert other.get(f"/a/{aid}/heatmap").status_code == 200
    # The resume code is shown once only.
    assert '<code class="token">' not in client.get(f"/a/{aid}").text


def test_retake_and_progress_page(app, client):
    root = SAMPLES[0].parent
    old = json.loads((root / "food_bank.json").read_text())
    new = json.loads((root / "retakes" / "food_bank.json").read_text())
    aid, first = start(client, old["organisation"]["name"], old["organisation"]["profile"])
    token = re.search(r'<code class="token">([^<]+)</code>', first.text).group(1)
    form = lambda answers: {f"q_{q}": v for q, v in answers.items()}  # noqa: E731
    assert client.post(f"/a/{aid}/s/govern", data=form(old["answers"])).status_code == 200
    assert client.get(f"/a/{aid}/progress", follow_redirects=False).status_code == 303  # nothing to compare yet

    page = client.get(f"/a/{aid}/retake").text
    assert 'name="has_office" checked' in page and 'name="remote_staff">' in page
    flags = {f: "on" for f, on in new["organisation"]["profile"].items() if on}
    r = client.post(f"/a/{aid}/retake", data={**flags, "copy": "on"})
    assert r.status_code == 200 and "Version 2 started" in r.text
    v2 = int(re.search(r"/a/(\d+)", str(r.url)).group(1))
    assert v2 != aid
    # Everything but the new remote-working questions is already answered.
    assert re.search(r"(\d+) of (\d+) questions answered", r.text).groups() == ("86", "88")

    changed = {q: v for q, v in new["answers"].items() if old["answers"].get(q) != v}
    client.post(f"/a/{v2}/s/govern", data=form(changed))
    progress = client.get(f"/a/{v2}/progress")
    assert progress.status_code == 200
    assert "went from 27% to 34% of its target profile" in progress.text
    assert "4 completed, 14 moved forward" in progress.text
    assert "remote staff (now yes)" in progress.text

    wb = load_workbook(io.BytesIO(client.get(f"/a/{v2}/workbook.xlsx").content))
    assert wb["Progress"]["B1"].value == "version 2 with version 1"
    doc = Document(io.BytesIO(client.get(f"/a/{v2}/report.docx").content))
    assert any(p.text == "Progress since version 1" for p in doc.paragraphs)

    # Version 1 is unchanged, and the one resume code opens both versions in another browser.
    assert "86 of 86 questions answered" in client.get(f"/a/{aid}").text
    with TestClient(app) as other:
        assert other.post("/resume", data={"assessment_id": aid, "token": token}).status_code == 200
        assert other.get(f"/a/{v2}/progress").status_code == 200
        assert "what changed" in other.get(f"/a/{aid}").text
