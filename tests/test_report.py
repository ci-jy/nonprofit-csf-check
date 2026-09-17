import subprocess
import sys

import pytest
from docx import Document

from csfcheck.report import ACTION_HEADERS, CATEGORY_HEADERS, PROFILE_HEADERS, TOP_ACTIONS, build_report
from csfcheck.samples import load_answers

from .conftest import ROOT, SAMPLES


def report_for(engine, path):
    org = load_answers(path)
    result = engine.assess(org.profile, org.answers)
    return org, result, build_report(engine, result, org.answers, org.evidence, org.name)


def cells(table):
    return [[c.text for c in row.cells] for row in table.rows]


@pytest.mark.parametrize("path", SAMPLES, ids=[p.stem for p in SAMPLES])
def test_report_sections_and_tables(engine, path):
    org, result, doc = report_for(engine, path)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings[:4] == ["Executive summary", "Profile by function", "Profile by category",
                            f"Top {TOP_ACTIONS} actions"]
    assert headings[-1] == "About this assessment"

    profile, categories, actions = doc.tables
    assert cells(profile)[0] == list(PROFILE_HEADERS)
    assert [r[0] for r in cells(profile)[1:]] == ["Govern", "Identify", "Protect", "Detect", "Respond", "Recover"]
    assessed = sum(int(r[1]) for r in cells(profile)[1:])
    assert assessed == result.totals["subcategories_assessed"]
    assert sum(int(r[2]) for r in cells(profile)[1:]) == result.totals["subcategories_at_target"]

    assert cells(categories)[0] == list(CATEGORY_HEADERS)
    assert len(categories.rows) == 1 + len(engine.catalog.categories)

    rows = cells(actions)
    assert rows[0] == list(ACTION_HEADERS)
    assert len(rows) == 1 + TOP_ACTIONS
    for row, action in zip(rows[1:], result.actions):
        assert row[:3] == [str(action.rank), action.subcategory, action.what_to_do]
        assert float(row[5]) == action.risk_reduction


@pytest.mark.parametrize("path", SAMPLES, ids=[p.stem for p in SAMPLES])
def test_executive_summary_matches_the_scores(engine, path):
    org, result, doc = report_for(engine, path)
    t = result.totals
    text = "\n".join(p.text for p in doc.paragraphs)
    assert f"{org.name} answered {t['questions_asked']} of the {t['questions_in_bank']} questions" in text
    assert f"Of the {t['subcategories_assessed']} CSF 2.0 subcategories that apply, " \
           f"{t['subcategories_at_target']} meet the target" in text
    assert f"The action plan has {len(result.actions)} actions" in text
    numbered = [p.text for p in doc.paragraphs if p.style.name == "List Number"]
    assert [n.split(" (")[0] for n in numbered] == [a.what_to_do for a in result.actions[:5]]
    assert "not an ISO 27001 or SOC 2 certification" in text
    assert "not finished" not in text


def test_not_applicable_section_gives_reasons(engine):
    org, result, doc = report_for(engine, ROOT / "samples" / "student_club.json")
    excluded = [r for r in result.subcategories if r.status == "excluded"]
    assert excluded
    text = "\n".join(p.text for p in doc.paragraphs)
    for r in excluded:
        assert r.id in text and r.reason in text


def test_incomplete_assessment_is_flagged(engine):
    org = load_answers(SAMPLES[0])
    partial = dict(list(org.answers.items())[:10])
    result = engine.assess(org.profile, partial)
    doc = build_report(engine, result, partial, {}, org.name)
    assert any("not finished" in p.text for p in doc.paragraphs)


def test_cli_report(tmp_path):
    out = tmp_path / "club.docx"
    proc = subprocess.run(
        [sys.executable, "-m", "csfcheck", "report", "samples/student_club.json", "-o", str(out), "--top", "5"],
        cwd=ROOT, capture_output=True, text=True,
    )
    assert proc.returncode == 0, proc.stderr
    assert "wrote" in proc.stdout
    assert len(Document(out).tables[2].rows) == 6
