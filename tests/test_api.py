import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from csfcheck.api import create_app
from csfcheck.engine import ASKED
from csfcheck.workbook import SHEETS

from .conftest import ROOT, SAMPLES


@pytest.fixture(scope="module")
def client(database_url):
    app = create_app(database_url)
    with TestClient(app) as c:
        yield c
    app.state.db.dispose()


def start(client, name, profile):
    org = client.post("/organisations", json={"name": name, "profile": profile})
    assert org.status_code == 201, org.text
    a = client.post(f"/organisations/{org.json()['id']}/assessments")
    assert a.status_code == 201
    body = a.json()
    return body["id"], {"X-Assessment-Token": body["token"]}


def test_health(client):
    body = client.get("/health").json()
    assert body["subcategories"] == 106 and body["questions"] == 100


@pytest.mark.parametrize("path", SAMPLES, ids=[p.stem for p in SAMPLES])
def test_sample_completes_through_api(client, engine, path):
    sample = json.loads(path.read_text())
    org = sample["organisation"]
    aid, headers = start(client, org["name"], org["profile"])
    skipped_by_profile = {q.id for q in engine.bank.questions if not q.applies_to(org["profile"])}
    assert skipped_by_profile, "every sample has some irrelevant questions"

    asked = []
    for _ in range(100):
        nxt = client.get(f"/assessments/{aid}/next?limit=6", headers=headers).json()
        if nxt["complete"]:
            break
        batch = [q["id"] for q in nxt["questions"]]
        assert batch and not set(batch) & skipped_by_profile
        asked += batch
        payload = [
            {"question_id": q, "value": sample["answers"][q], "evidence": sample["evidence"].get(q, "")}
            for q in batch
        ]
        r = client.post(f"/assessments/{aid}/answers", json={"answers": payload}, headers=headers)
        assert r.status_code == 200, r.text
    assert len(asked) == len(set(asked))
    assert set(asked) == set(sample["answers"])  # the API asks exactly what the sample answers

    profile = client.get(f"/assessments/{aid}/profile", headers=headers).json()
    assert profile["complete"]
    assert len(profile["subcategories"]) == 106
    totals = profile["totals"]
    assert totals["questions_asked"] == len(asked)
    assert totals["skipped_profile"] == len(skipped_by_profile)

    plan = client.get(f"/assessments/{aid}/action-plan", headers=headers).json()["actions"]
    assert plan, "every sample has gaps to close"
    assert [a["rank"] for a in plan] == list(range(1, len(plan) + 1))
    priorities = [a["priority"] for a in plan]
    assert priorities == sorted(priorities, reverse=True)

    wb_resp = client.get(f"/assessments/{aid}/workbook", headers=headers)
    assert wb_resp.status_code == 200
    wb = load_workbook(io.BytesIO(wb_resp.content))
    assert wb.sheetnames == list(SHEETS)
    assert wb["Action Plan"].max_row - 1 == len(plan)
    assert wb["Evidence Log"].max_row - 1 == len(asked)
    evidence = {row[0].value: row[6].value for row in wb["Evidence Log"].iter_rows(min_row=2)}
    for qid, note in sample["evidence"].items():
        assert evidence[qid] == note


def test_skipped_question_answer_is_rejected(client):
    aid, headers = start(client, "Tiny club", {})
    r = client.post(
        f"/assessments/{aid}/answers",
        json={"answers": [{"question_id": "server-health", "value": "yes"}]},
        headers=headers,
    )
    assert r.status_code == 422 and "skipped" in r.json()["detail"]
    r = client.post(
        f"/assessments/{aid}/answers",
        json={"answers": [{"question_id": "gv-owner", "value": "maybe"}]},
        headers=headers,
    )
    assert r.status_code == 422


def test_lowering_a_parent_drops_its_follow_up(client, engine):
    aid, headers = start(client, "Backup test", {})
    post = lambda answers: client.post(  # noqa: E731
        f"/assessments/{aid}/answers",
        json={"answers": [{"question_id": q, "value": v} for q, v in answers.items()]},
        headers=headers,
    )
    assert post({"pr-backup": "yes", "pr-backup-test": "yes"}).json()["answered"] == 2
    assert post({"pr-backup": "no"}).json()["answered"] == 1
    profile = client.get(f"/assessments/{aid}/profile", headers=headers).json()
    rp03 = next(s for s in profile["subcategories"] if s["id"] == "RC.RP-03")
    assert rp03["current"] == 0


def test_token_is_required(client):
    aid, headers = start(client, "Token check", {"has_office": True})
    assert client.get(f"/assessments/{aid}/next").status_code == 422
    bad = {"X-Assessment-Token": "not-the-token"}
    assert client.get(f"/assessments/{aid}/next", headers=bad).status_code == 404
    assert client.get(f"/assessments/{aid}/next", headers=headers).status_code == 200


def test_unknown_profile_flag_is_rejected(client):
    r = client.post("/organisations", json={"name": "x", "profile": {"has_pets": True}})
    assert r.status_code == 422


def test_retake_keeps_the_token_and_compares_versions(client, engine):
    old = json.loads((ROOT / "samples" / "food_bank.json").read_text())
    new = json.loads((ROOT / "samples" / "retakes" / "food_bank.json").read_text())
    aid, headers = start(client, old["organisation"]["name"], old["organisation"]["profile"])
    answer = lambda a, answers: client.post(  # noqa: E731
        f"/assessments/{a}/answers", headers=headers,
        json={"answers": [{"question_id": q, "value": v} for q, v in answers.items()]},
    )
    assert answer(aid, old["answers"]).status_code == 200
    assert client.get(f"/assessments/{aid}/compare", headers=headers).status_code == 404  # nothing earlier

    r = client.post(f"/assessments/{aid}/retake", headers=headers, json={"profile": new["organisation"]["profile"]})
    assert r.status_code == 201, r.text
    v2 = r.json()
    assert v2["version"] == 2 and v2["previous_id"] == aid
    assert v2["answers_copied"] == len(old["answers"])  # same token, answers carried over
    nxt = client.get(f"/assessments/{v2['id']}/next?limit=100", headers=headers).json()
    # Only the questions the new profile adds; follow-ups wait for their parents to improve.
    assert {q["id"] for q in nxt["questions"]} == {"remote-devices", "remote-lost-device"}

    changed = {q: v for q, v in new["answers"].items() if old["answers"].get(q) != v}
    parents = {q: v for q, v in changed.items() if not engine.bank.by_id[q].requires}
    assert answer(v2["id"], parents).status_code == 200
    assert answer(v2["id"], {q: v for q, v in changed.items() if q not in parents}).status_code == 200
    assert client.get(f"/assessments/{v2['id']}/next", headers=headers).json()["complete"]

    # The first version is still scored on its own profile and answers.
    first = client.get(f"/assessments/{aid}/profile", headers=headers).json()
    assert first["version"] == 1 and first["totals"]["subcategories_at_target"] == 19

    diff = client.get(f"/assessments/{v2['id']}/compare", headers=headers).json()
    assert diff["before"] == {"id": aid, "version": 1}
    expected = engine.assess(new["organisation"]["profile"], new["answers"])
    assert diff["totals"]["at_target_after"] == expected.totals["subcategories_at_target"] == 22
    assert (diff["totals"]["improved"], diff["totals"]["regressed"], diff["totals"]["actions_completed"]) == (18, 3, 4)
    assert diff["profile_changes"] == {"remote_staff": {"before": False, "after": True}}

    listed = client.get(f"/assessments/{aid}/versions", headers=headers).json()["versions"]
    assert [v["version"] for v in listed] == [1, 2]

    wb = load_workbook(io.BytesIO(client.get(f"/assessments/{v2['id']}/workbook", headers=headers).content))
    summary = {row[0]: row[1] for row in wb["Progress"].iter_rows(max_row=13, values_only=True)}
    assert summary["Compared"] == "version 2 with version 1"
    assert summary["Actions completed"] == "4 of 85"


def test_compare_refuses_another_organisations_assessment(client):
    a1, h1 = start(client, "Club A", {})
    a2, h2 = start(client, "Club B", {})
    assert client.get(f"/assessments/{a1}/compare?with={a2}", headers=h1).status_code == 404
    assert client.post(f"/assessments/{a1}/retake", headers=h2).status_code == 404
    v2 = client.post(f"/assessments/{a1}/retake", headers=h1).json()
    assert client.get(f"/assessments/{v2['id']}/compare?with={a1}", headers=h1).status_code == 200
    assert client.get(f"/assessments/{v2['id']}/compare", headers=h2).status_code == 404
