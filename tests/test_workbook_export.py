import json
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from csfcheck.workbook import (
    ACTION_HEADERS,
    EVIDENCE_HEADERS,
    GAP_HEADERS,
    PROFILE_HEADERS,
    PROGRESS_ACTION_HEADERS,
    PROGRESS_ANSWER_HEADERS,
    PROGRESS_FUNCTION_HEADERS,
    PROGRESS_SUBCATEGORY_HEADERS,
    SHEETS,
)

from .conftest import ROOT, SAMPLES


@pytest.mark.parametrize("path", SAMPLES, ids=[p.stem for p in SAMPLES])
def test_cli_export_writes_verified_sheets(tmp_path, engine, path):
    out = tmp_path / f"{path.stem}.xlsx"
    proc = subprocess.run(
        [sys.executable, "-m", "csfcheck", "export", str(path), "-o", str(out)],
        cwd=ROOT, capture_output=True, text=True,
    )
    assert proc.returncode == 0, proc.stderr
    sample = json.loads(path.read_text())
    result = engine.assess(sample["organisation"]["profile"], sample["answers"])

    wb = load_workbook(out)
    assert wb.sheetnames == list(SHEETS)
    headers = lambda ws: [c.value for c in ws[1]]  # noqa: E731
    assert headers(wb["Profile"]) == list(PROFILE_HEADERS)
    assert headers(wb["Gaps"]) == list(GAP_HEADERS)
    assert headers(wb["Action Plan"]) == list(ACTION_HEADERS)
    assert headers(wb["Evidence Log"]) == list(EVIDENCE_HEADERS)

    profile_rows = list(wb["Profile"].iter_rows(min_row=2, values_only=True))
    assert len(profile_rows) == 106
    statuses = {r[2]: r[4] for r in profile_rows}
    assert sum(s == "Not applicable" for s in statuses.values()) == result.totals["subcategories_excluded"]

    gaps = list(wb["Gaps"].iter_rows(min_row=2, values_only=True))
    assert len(gaps) == 106 - result.totals["subcategories_excluded"] - result.totals["subcategories_at_target"]
    risks = [g[8] for g in gaps]
    assert risks == sorted(risks, reverse=True)

    plan = list(wb["Action Plan"].iter_rows(min_row=2, values_only=True))
    assert [p[0] for p in plan] == list(range(1, len(result.actions) + 1))
    assert {p[2] for p in plan} == {g[0] for g in gaps}  # one action per gap

    log = list(wb["Evidence Log"].iter_rows(min_row=2, values_only=True))
    assert {r[0] for r in log} == set(sample["answers"])
    assert all(r[3] != "(unanswered)" for r in log)


def test_first_assessment_progress_sheet_is_a_baseline(tmp_path, engine):
    out = tmp_path / "club.xlsx"
    subprocess.run([sys.executable, "-m", "csfcheck", "export", "samples/student_club.json", "-o", str(out)],
                   cwd=ROOT, check=True, capture_output=True)
    rows = list(load_workbook(out)["Progress"].iter_rows(values_only=True))
    assert "first assessment" in rows[0][1]
    assert [r[0] for r in rows[-6:]] == ["Govern", "Identify", "Protect", "Detect", "Respond", "Recover"]


def test_export_with_previous_writes_the_progress_sheet(tmp_path):
    out = tmp_path / "food_bank_v2.xlsx"
    proc = subprocess.run(
        [sys.executable, "-m", "csfcheck", "export", "samples/retakes/food_bank.json",
         "--previous", "samples/food_bank.json", "-o", str(out)],
        cwd=ROOT, capture_output=True, text=True,
    )
    assert proc.returncode == 0, proc.stderr
    assert "18 subcategories improved, 3 regressed, 4 of 85 actions completed" in proc.stdout
    ws = load_workbook(out)["Progress"]
    rows = list(ws.iter_rows(values_only=True))
    summary = {r[0]: r[1] for r in rows[:13]}
    assert summary["Subcategories improved"] == 18 and summary["Subcategories regressed"] == 3
    assert summary["Profile changes"] == "remote_staff: yes (was no)"

    def table(header):
        start = next(i for i, r in enumerate(rows) if r[: len(header)] == tuple(header))
        body = []
        for r in rows[start + 1:]:
            if r[0] is None:
                break
            body.append(r)
        return body

    functions = {r[0]: r for r in table(PROGRESS_FUNCTION_HEADERS)}
    assert functions["Respond"][1:4] == (10, 32, 22)
    subs = table(PROGRESS_SUBCATEGORY_HEADERS)
    assert len(subs) == 21 and {r[3] for r in subs} == {"improved", "regressed"}
    actions = table(PROGRESS_ACTION_HEADERS)
    assert [r[2] for r in actions].count("completed") == 4
    assert [r[2] for r in actions][:4] == ["completed"] * 4  # done ones first
    answers = table(PROGRESS_ANSWER_HEADERS)
    assert len(answers) == 17


def test_export_rejects_a_bad_previous_file(tmp_path):
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"organisation": {"name": "x", "profile": {}}, "answers": {"gv-owner": "maybe"}}))
    proc = subprocess.run(
        [sys.executable, "-m", "csfcheck", "export", "samples/food_bank.json", "--previous", str(bad),
         "-o", str(tmp_path / "x.xlsx")],
        cwd=ROOT, capture_output=True, text=True,
    )
    assert proc.returncode == 2 and "bad.json" in proc.stderr
